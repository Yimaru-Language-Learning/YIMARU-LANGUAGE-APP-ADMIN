#!/usr/bin/env python3
"""Generate Yimaru Admin Panel Developer Documentation (.docx)."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "Yimaru_Admin_Developer_Documentation.docx"


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_para(
    doc,
    text,
    *,
    size=11,
    bold=False,
    italic=False,
    space_after=8,
    space_before=0,
    align=None,
):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading_custom(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(
            run,
            name="Calibri",
            size={1: 18, 2: 14, 3: 12}.get(level, 11),
            bold=True,
        )
        if level == 1:
            run.font.color.rgb = RGBColor(0x5E, 0x1F, 0x70)
        elif level == 2:
            run.font.color.rgb = RGBColor(0x9E, 0x28, 0x91)
    return h


def add_bullets(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run_font(run, size=size)


def add_numbered(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run_font(run, size=size)


def add_code_block(doc, text):
    for line in text.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Inches(0.15)
        run = p.add_run(line if line else " ")
        set_run_font(run, name="Consolas", size=9, color=RGBColor(0x33, 0x33, 0x33))
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=10, bold=True, color=RGBColor(0x5E, 0x1F, 0x70))
    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = ""
            p = cells[c_i].paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(10)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def build():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    # Title
    for _ in range(3):
        doc.add_paragraph()
    add_para(doc, "Yimaru Admin Panel", size=28, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    add_para(doc, "Developer Documentation", size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
    add_para(
        doc,
        "A practical guide for engineers building and maintaining the Yimaru Academy admin SPA",
        size=12,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=28,
    )
    add_para(
        doc,
        f"Package: yimaru-admin (React 19 · TypeScript · Vite 7)\n"
        f"Document generated {date.today().isoformat()}\n"
        f"Audience: Frontend / full-stack developers working on the admin console\n"
        f"Source: https://gitea.yaltopia.com/Yimaru/Yimaru-Admin.git",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )
    add_para(
        doc,
        "Internal / proprietary — Yaltopia Tech · Yimaru Academy",
        size=9,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    doc.add_page_break()

    # 1 Intro
    add_heading_custom(doc, "1. Introduction", 1)
    add_para(
        doc,
        "The Yimaru Admin Panel is the staff-facing web app for Yimaru Academy. Content "
        "managers reorder courses at midnight, support agents triage issues, finance peeks at "
        "revenue charts, and super-admins invite the next teammate—all through this SPA. It does "
        "not own business rules; it talks to Yimaru-Backend over HTTPS and presents them safely.",
    )
    add_para(
        doc,
        "This document is written for developers who will change routes, API clients, permission "
        "gates, or UI. It assumes React and TypeScript familiarity. It does not assume you already "
        "know why we bootstrap the team session before first paint, or why invite payloads must "
        "send SUPER_ADMIN instead of the label “Super Admin.”",
    )
    add_para(doc, "Keep these siblings open while you work:", space_after=4)
    add_bullets(
        doc,
        [
            "Backend: https://gitea.yaltopia.com/Yimaru/Yimaru-BackEnd.git",
            "Backend docs/*-admin-panel-integration.md — API contracts for invites, exports, practices, etc.",
            "This repo’s README.md — quick start and env tables",
            "Backend Swagger on a running API — /swagger/index.html",
        ],
    )

    add_heading_custom(doc, "1.1 What the admin app is (and is not)", 2)
    add_para(
        doc,
        "It is a team-member console: authentication is against team_members, not learner users. "
        "Permissions come from RBAC keys on the JWT / member payload. It is not the mobile learner "
        "app, and it should not invent entitlement logic that the API does not enforce.",
    )

    add_heading_custom(doc, "1.2 Branches", 2)
    add_para(
        doc,
        "We keep main and production aligned for releases. Typical flow: land on main, push, "
        "fast-forward production, push production. If only one branch is updated, deploys drift—"
        "and someone will ask why UAT looks different from production.",
    )

    # 2 Stack
    add_heading_custom(doc, "2. Tech stack and runtime shape", 1)
    add_table(
        doc,
        ["Concern", "Choice"],
        [
            ["UI library", "React 19 + React DOM"],
            ["Language", "TypeScript ~5.9 (strict-ish project defaults)"],
            ["Bundler / dev server", "Vite 7"],
            ["Routing", "react-router-dom 7 — see src/app/AppRoutes.tsx"],
            ["HTTP", "Axios instance in src/api/http.ts"],
            ["Styling", "Tailwind 3 + CSS variables in index.css"],
            ["Primitives", "Radix UI (dialog, dropdown, avatar, …)"],
            ["Icons", "lucide-react"],
            ["Charts", "Recharts"],
            ["Drag and drop", "@dnd-kit (reorder pages)"],
            ["Toasts", "Sonner"],
            ["Global state", "Zustand (src/zustand) + lots of local React state"],
            ["Uploads", "tus-js-client where resumable/Vimeo flows need it"],
        ],
    )
    add_para(
        doc,
        "Build metadata: vite.config.ts injects __BUILD_HASH__ (git short SHA) and __BUILD_TIME__. "
        "Handy when support asks “which admin build is live?”",
        space_before=6,
    )

    # 3 Architecture
    add_heading_custom(doc, "3. Application architecture", 1)
    add_para(
        doc,
        "Think in layers that stay boring on purpose:",
    )
    add_code_block(
        doc,
        """
main.tsx → App.tsx (session bootstrap) → AppRoutes
  └─ public auth/legal routes
  └─ AppLayout (Sidebar + Topbar + <Outlet />)
       └─ page components under src/pages/*
            └─ call src/api/*.ts
                 └─ Axios http.ts (Bearer + refresh queue)
""",
    )
    add_table(
        doc,
        ["Folder", "Put here…"],
        [
            ["src/pages/", "Route screens and feature-specific UI"],
            ["src/components/", "Reusable UI (ui/, sidebar/, analytics/, dashboard/)"],
            ["src/api/", "One module per backend domain"],
            ["src/types/", "TS shapes mirroring API JSON (usually snake_case fields)"],
            ["src/lib/", "Pure helpers: auth storage, teamRoles, analytics formatters, errors"],
            ["src/hooks/", "Shared hooks (e.g. useTeamPermissions)"],
            ["src/layouts/", "AppLayout shell"],
            ["src/zustand/", "Cross-page client state when props drilling hurts"],
        ],
    )
    add_para(
        doc,
        "Canonical routing truth is AppRoutes.tsx—not the sidebar labels, not this document if "
        "they diverge. When you add a page, register the route and wire nav (with permission) in "
        "Sidebar.tsx.",
        space_before=6,
    )

    # 4 Setup
    add_heading_custom(doc, "4. Local development setup", 1)
    add_heading_custom(doc, "4.1 Prerequisites", 2)
    add_bullets(
        doc,
        [
            "Node.js 18+ (20 LTS recommended)",
            "npm (package-lock.json is the source of truth)",
            "A running Yimaru Backend reachable from the browser",
            "A team account with the permissions you need to test",
        ],
    )
    add_heading_custom(doc, "4.2 First run", 2)
    add_numbered(
        doc,
        [
            "git clone https://gitea.yaltopia.com/Yimaru/Yimaru-Admin.git && cd Yimaru-Admin",
            "npm install",
            "Create .env with VITE_API_BASE_URL (and VITE_GOOGLE_CLIENT_ID if testing Google login)",
            "npm run dev — usually http://localhost:5173",
            "Log in with a team member; confirm Network tab hits your API prefix",
        ],
    )
    add_para(
        doc,
        "Example env (adjust host/port to your backend):",
        space_before=6,
        space_after=4,
    )
    add_code_block(
        doc,
        """
VITE_API_BASE_URL=http://localhost:8080/api/v1
VITE_GOOGLE_CLIENT_ID=your-client-id.apps.googleusercontent.com
""",
    )
    add_para(
        doc,
        "Vite only exposes VITE_* variables, and they are baked in at build time for production "
        "bundles. Changing .env requires restarting the dev server; changing production API URL "
        "requires a rebuild.",
    )

    add_heading_custom(doc, "4.3 Scripts", 2)
    add_table(
        doc,
        ["Script", "Purpose"],
        [
            ["npm run dev", "Vite HMR"],
            ["npm run build", "Production bundle to dist/"],
            ["npm run preview", "Serve dist locally"],
            ["npm run lint", "ESLint"],
        ],
    )

    # 5 Auth
    add_heading_custom(doc, "5. Authentication and session lifecycle", 1)
    add_para(
        doc,
        "Admin auth is team-based. Learners cannot log into this app with a student password and "
        "should not be able to—if they can, something is very wrong on the backend route guards.",
    )
    add_heading_custom(doc, "5.1 Public vs shell routes", 2)
    add_table(
        doc,
        ["Path", "Purpose"],
        [
            ["/login", "Email/password (+ Google when configured)"],
            ["/forgot-password", "Request reset email"],
            ["/reset-password", "Complete reset with OTP from link"],
            ["/verification", "OTP UI where the flow still uses it"],
            ["/accept-invite", "Invitee sets profile/password"],
            ["/about, /terms, /privacy, /account-deletion", "Public/legal"],
            ["Everything under AppLayout", "Requires a usable team session"],
        ],
    )
    add_heading_custom(doc, "5.2 Storage and http.ts", 2)
    add_para(
        doc,
        "Tokens and member metadata live in browser storage via src/lib/teamAuthStorage "
        "(access_token, refresh_token, member_id, team_role, …). The shared Axios client:",
    )
    add_bullets(
        doc,
        [
            "Sets Authorization: Bearer <access_token> on outbound calls",
            "Proactively refreshes when the access JWT is near expiry",
            "Queues concurrent 401s while a single /team/refresh is in flight",
            "Clears session and redirects to /login on hard failures (including inactive account query flag)",
            "Can skipErrorToast for flows that handle errors inline (invites, some auth)",
        ],
    )
    add_para(
        doc,
        "App.tsx bootstraps the browser session before the first protected render. That exists "
        "because we used to flash “Authorization header missing” toasts when clearTeamSession "
        "raced with early API calls. Do not “simplify” bootstrap away without re-testing cold loads.",
        space_before=6,
    )

    add_heading_custom(doc, "5.3 Team invites — role keys and email checks", 2)
    add_para(
        doc,
        "POST /team/members/invite expects team_role as a built-in key (SUPER_ADMIN), an exact "
        "RBAC role name, or a numeric roles.id. Display labels like “Super Admin” are for humans. "
        "Helpers in src/lib/teamRoles.ts (teamRoleNameForInvite, rbacRolesToTeamRoleOptions) keep "
        "the invite dialog honest. The dialog also loads existing team emails and blocks "
        "duplicates before send—invites cannot target an email that already exists on team_members.",
    )

    # 6 Permissions
    add_heading_custom(doc, "6. Permissions (RBAC) in the UI", 1)
    add_para(
        doc,
        "Gate features on permission keys, not on team_role === \"ADMIN\". Sidebar entries use "
        "useTeamPermissions and helpers like hasFaqPermission, hasRatingsPermission, "
        "hasExportPermission, etc. fetchCurrentTeamMemberPermissions in team.api.ts walks the "
        "members list to find the signed-in member’s permissions array (and syncs team_role).",
    )
    add_para(
        doc,
        "When you add a new admin action, ask backend for the permission string they wired in "
        "RequirePermission, then gate the button and ideally the route. Hiding a nav link is not "
        "security—the API must still refuse—but showing a dead-end page is a bad UX.",
    )

    # 7 Navigation map
    add_heading_custom(doc, "7. Product areas and routing map", 1)
    add_para(
        doc,
        "High-level navigation (see Sidebar.tsx + AppRoutes.tsx for exact paths):",
    )
    add_table(
        doc,
        ["Area", "Typical paths", "Notes"],
        [
            ["Dashboard", "/dashboard", "KPIs, charts, plans peek"],
            ["Analytics", "/analytics", "Tabbed deep analytics"],
            ["Users", "/users/list, /users/:id, groups, deletion-requests", "Learner admin"],
            ["Roles", "/roles, /roles/add", "RBAC + invite-into-role"],
            ["Team", "/team, /team/add, /team/:id", "Staff directory"],
            ["New content", "/new-content, /reorder, question-types, courses…", "Primary CMS"],
            ["Legacy content", "/content/*", "Older category flows; still routed"],
            ["Personas", "/personas", "permission: personas.list"],
            ["Notifications", "/notifications/* + email templates", "Create/schedule"],
            ["Payments", "/payments", "Payment ops"],
            ["Subscriptions export", "subscriptions export routes", "CSV helpers"],
            ["Issues", "/issues", "Use status rejected, not closed"],
            ["FAQs / ratings / user-log", "/faqs, ratings, /user-log", "Permission gated"],
            ["Profile / settings", "/profile, /settings", "Self-service team profile"],
        ],
    )

    add_heading_custom(doc, "7.1 New content vs legacy /content", 2)
    add_para(
        doc,
        "Day-to-day CMS work happens under /new-content (Learn English programs and "
        "Duolingo/IELTS catalog, practices, question-type library, reorder). /content/* remains "
        "for older hierarchies. Prefer extending new-content unless you are fixing a legacy path "
        "that production still bookmarks.",
    )

    add_heading_custom(doc, "7.2 Reorder Content", 2)
    add_para(
        doc,
        "Route: /new-content/reorder. Tabs: Learn English | Duolingo/IELTS. Drag-and-drop via "
        "@dnd-kit; persist on drop through reorder* helpers in courses.api.ts. Shared list UI "
        "lives in ContentHierarchyList / ExamPrepContentHierarchyList. Keep the layout dense—"
        "editors reorder long trees and hate giant paddings.",
    )

    # 8 API layer
    add_heading_custom(doc, "8. API client layer", 1)
    add_para(
        doc,
        "Convention: one file per backend domain under src/api/. Pages import functions; they "
        "should not construct raw axios instances. Types belong in src/types/.",
    )
    add_table(
        doc,
        ["Module", "Domain"],
        [
            ["http.ts", "Shared client + interceptors"],
            ["team.api.ts / teamRefresh.api.ts", "Team auth, members, invites"],
            ["analytics.api.ts", "Dashboard analytics"],
            ["courses.api.ts", "LMS + exam-prep content & reorder"],
            ["users.api.ts", "Learners"],
            ["rbac.api.ts", "Roles & permissions"],
            ["notifications.api.ts / emailTemplates.api.ts", "Messaging"],
            ["payments.api.ts / subscription-plans.api.ts / admin-subscriptions.api.ts", "Money"],
            ["issues.api.ts / faq.api.ts / ratings.api.ts / personas.api.ts", "Support content"],
            ["files.api.ts", "Uploads"],
            ["activity-logs.api.ts", "Audit trail"],
            ["questionTypeDefinitions.api.ts / …Groups…", "Dynamic question types"],
            ["app-versions.api.ts", "Mobile version gates"],
            ["progress.api.ts", "Progress-related admin views if used"],
        ],
    )
    add_para(
        doc,
        "Errors: prefer notifyApiError / getApiErrorMessage from src/lib/apiErrors.ts. For "
        "multi-step dialogs (bulk invite), collect per-email outcomes instead of a single toast.",
        space_before=6,
    )
    add_para(
        doc,
        "When the backend publishes a new admin-panel-integration.md, update types + api module "
        "in the same PR as the UI. Drift between Swagger and our TypeScript types is how silent "
        "undefined bugs are born.",
    )

    # 9 Design system
    add_heading_custom(doc, "9. Design system and UI patterns", 1)
    add_para(
        doc,
        "Brand purple is intentional (brand-500 ≈ #9E2891). Gold and mint are accents for "
        "highlights and positive trends. Neutrals use grayScale-* backed by CSS variables "
        "(--gs-*). Font: Inter. Prefer components under src/components/ui/ over one-off markup.",
    )
    add_heading_custom(doc, "9.1 Dashboard / Analytics visual language", 2)
    add_para(
        doc,
        "Shared primitives (where present) and page patterns aim for: light page canvas, white "
        "cards with thin borders, underline tabs in brand color, KPI cards with label + thin "
        "icon + large value + trend pill, purple area charts with soft fills. Sensitive money "
        "fields use SensitiveValue / SensitiveRevealProvider so operators can hide revenue in "
        "screen shares.",
    )
    add_heading_custom(doc, "9.2 Accessibility and density", 2)
    add_para(
        doc,
        "Radix primitives give us focus management for dialogs/menus—keep using them. For dense "
        "ops screens (reorder, issue queues), favor information density over marketing whitespace, "
        "but never remove hit targets for drag handles and primary actions.",
    )

    # 10 Feature deep dives
    add_heading_custom(doc, "10. Feature deep dives", 1)

    add_heading_custom(doc, "10.1 Analytics", 2)
    add_para(
        doc,
        "Data comes from getDashboard with DashboardFilters (all_time / year / year_month / "
        "custom). AnalyticsTimeRangeFilter owns the date UX. Tabs conceptually split Learn "
        "English, Content Performance, and Subscription & Revenue so operators are not scrolling "
        "a single endless page. Formatters and pie builders live in src/lib/analytics.ts—"
        "including label cleanup (Unknown → Other) to match backend normalization.",
    )

    add_heading_custom(doc, "10.2 Users and subscriptions on user detail", 2)
    add_para(
        doc,
        "User detail composes profile, learning activity, and subscription admin "
        "(grant/extend/cancel) against admin-subscriptions APIs. Follow backend integration docs "
        "for payload shapes; do not guess lifecycle transitions.",
    )

    add_heading_custom(doc, "10.3 Roles and bulk deactivate", 2)
    add_para(
        doc,
        "RolesListPage can invite into a role, edit permissions, and bulk deactivate/reactivate. "
        "SUPER_ADMIN guardrails are enforced server-side—surface API errors clearly instead of "
        "retrying blindly.",
    )

    add_heading_custom(doc, "10.4 Issues", 2)
    add_para(
        doc,
        "Valid PATCH statuses include pending, in_progress, resolved, rejected. The UI once sent "
        "closed, which failed validation. Keep rejected in the selector; legacy closed may still "
        "display for old rows.",
    )

    add_heading_custom(doc, "10.5 Notifications & email templates", 2)
    add_para(
        doc,
        "Creating notifications supports audiences by user ids, platform roles, team members, "
        "team roles, and direct emails (depending on channel). Scheduled sends need RFC3339 "
        "timestamps the API accepts. Email templates are a separate CRUD surface—test with "
        "Resend configured on the backend.",
    )

    # 11 Shipping
    add_heading_custom(doc, "11. How we like to ship UI changes", 1)
    add_numbered(
        doc,
        [
            "Confirm the API exists (or is shipping in the same release) and note the permission key.",
            "Add/extend types in src/types and functions in src/api.",
            "Build the page or dialog with existing ui/ components.",
            "Gate nav and destructive actions with permissions.",
            "Handle loading, empty, and error states—ops users hate infinite spinners with no retry.",
            "npm run lint; smoke the happy path and one failure path (403, validation).",
            "Commit with why; push main and sync production when releasing.",
        ],
    )
    add_para(
        doc,
        "Avoid drive-by refactors in the same PR as a production hotfix. Match local naming and "
        "file layout. Prefer snake_case field names in API types to match JSON rather than "
        "renaming everything to camelCase at the boundary—unless a module already maps explicitly.",
        space_before=6,
    )

    # 12 Deploy
    add_heading_custom(doc, "12. Build and deploy notes", 1)
    add_code_block(
        doc,
        """
npm run lint
npm run build
# deploy dist/ behind HTTPS; set VITE_* at build time
""",
    )
    add_para(
        doc,
        "Coordinate TEAM_INVITE_BASE_URL / TEAM_PASSWORD_RESET_BASE_URL on the backend with the "
        "deployed admin origin so invite and reset emails deep-link correctly.",
    )

    # 13 Troubleshooting
    add_heading_custom(doc, "13. Troubleshooting field guide", 1)
    add_table(
        doc,
        ["Symptom", "Likely cause / fix"],
        [
            ["CORS / network errors", "Wrong VITE_API_BASE_URL; backend CORS; mixed content"],
            ["Immediate bounce to login", "Refresh failing; cleared storage; JWT_KEY rotated on API"],
            ["401 toasts on first paint", "Session bootstrap race—keep App.tsx + http.ts behavior"],
            ["Invite invalid team role", "Sending display label; use SUPER_ADMIN / RBAC name / id"],
            ["Invite email already exists", "Member or pending invite exists; dialog should block"],
            ["Issue status PATCH fails", "Do not send closed; use rejected"],
            ["Empty analytics", "API error, date filter, or missing permission"],
            ["Nav item missing", "Permission helper returned false; check role permissions"],
            ["Google login broken", "VITE_GOOGLE_CLIENT_ID + backend Google OAuth env"],
            ["Production looks stale", "main pushed but production not updated (or CDN cache)"],
        ],
    )

    # 14 Glossary
    add_heading_custom(doc, "14. Glossary", 1)
    add_table(
        doc,
        ["Term", "Meaning in this app"],
        [
            ["Team member", "Staff account in team_members; admin login identity"],
            ["Permission key", "String like team.members.invite used to show/hide actions"],
            ["New content", "Primary CMS under /new-content"],
            ["Duolingo/IELTS tab", "Exam-prep hierarchy in reorder UI"],
            ["SensitiveValue", "UI control to mask revenue/PII-ish numbers"],
            ["Accept invite", "Public page completing staff onboarding from email link"],
        ],
    )

    # 15 Related docs
    add_heading_custom(doc, "15. Related documentation", 1)
    add_bullets(
        doc,
        [
            "This repo README.md",
            "docs/course-management-api-integration.md (if present in docs/)",
            "Backend README + docs/Yimaru_Backend_Developer_Documentation.docx",
            "Backend docs/team-member-invitation-admin-panel-integration.md",
            "Backend docs/team-member-refresh-token-admin-panel-integration.md",
            "Backend docs/csv-export-admin-panel-integration.md",
            "Backend docs/bulk-scheduled-notifications-integration.md",
            "Backend practice/question-type/vimeo/file-upload integration guides",
        ],
    )

    # Closing
    add_heading_custom(doc, "16. Closing notes", 1)
    add_para(
        doc,
        "The admin panel is where humans operate a live learning business. Broken toasts, wrong "
        "role payloads, and ungated delete buttons create real support load. Prefer clear empty "
        "states, honest error messages, and permission-aware UI. When the API contract changes, "
        "update types and the screen together.",
    )
    add_para(
        doc,
        "If this document disagrees with AppRoutes.tsx or http.ts, trust the code—then fix the "
        "doc (or regenerate it from docs/generate_admin_developer_docx.py).",
        italic=True,
    )
    add_para(
        doc,
        "— Yaltopia engineering, for Yimaru Academy admin developers",
        space_before=16,
        italic=True,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
